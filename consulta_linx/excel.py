import time
from datetime import datetime, timedelta
from pathlib import Path

import pandas as pd
from openpyxl.utils import get_column_letter

from . import config
from .utils import log


def resumo_por_canal(result: pd.DataFrame) -> pd.DataFrame:
    return result.groupby("CANAL", as_index=False).agg(
        LINHAS=("CANAL", "size"),
        VALOR_PEDIDO=("VALOR_PEDIDO", "sum"),
        VALOR_NF=("VALOR_NF", "sum"),
    )


def salvar_excel(
    result: pd.DataFrame,
    destino: Path,
    inicio,
    fim,
    compat: bool,
) -> Path:
    destino.parent.mkdir(parents=True, exist_ok=True)
    temporary_file = destino.with_name(destino.stem + ".tmp.xlsx")
    with pd.ExcelWriter(
        temporary_file,
        engine="openpyxl",
        datetime_format="dd/mm/yyyy",
        date_format="dd/mm/yyyy",
    ) as writer:
        result.to_excel(writer, sheet_name="Pedidos", index=False)
        summary = resumo_por_canal(result)
        summary.to_excel(writer, sheet_name="Resumo", index=False)
        info = pd.DataFrame(
            {
                "PARÂMETRO": [
                    "Gerado em",
                    "Pedidos/notas a partir de",
                    "Até (inclusive)",
                    "Modo compatível",
                ],
                "VALOR": [
                    datetime.now().strftime("%d/%m/%Y %H:%M"),
                    inicio.strftime("%d/%m/%Y"),
                    (fim - timedelta(days=1)).strftime("%d/%m/%Y"),
                    "Sim" if compat else "Não",
                ],
            }
        )
        info.to_excel(writer, sheet_name="Resumo", index=False, startrow=len(summary) + 3)

        worksheet = writer.sheets["Pedidos"]
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = (
            f"A1:{get_column_letter(len(config.COLUNAS_SAIDA))}{max(len(result) + 1, 1)}"
        )
        widths = {
            "CANAL": 10,
            "PEDIDO": 12,
            "DATA_PEDIDO": 12,
            "TIPO_PEDIDO": 14,
            "DESCRICAO_TIPO_PEDIDO": 28,
            "CLASSIFICACAO_PEDIDO": 20,
            "STATUS_APROVACAO": 16,
            "STATUS_PEDIDO": 13,
            "DATA_FATURAMENTO": 14,
            "DATA_HORA_EMISSAO": 19,
            "DATA_HORA_REGISTRO": 19,
            "COD_VENDEDOR": 12,
            "VENDEDOR": 28,
            "COD_CLIENTE": 13,
            "CLIENTE": 40,
            "CIDADE": 22,
            "ESTADO": 7,
            "NF": 10,
            "COD_PRODUTO": 14,
            "PRODUTO": 45,
            "QUANTIDADE": 11,
            "VALOR_PEDIDO": 14,
            "VALOR_NF": 14,
        }
        for index, column in enumerate(config.COLUNAS_SAIDA, start=1):
            letter = get_column_letter(index)
            worksheet.column_dimensions[letter].width = widths.get(column, 12)
            number_format = (
                "#,##0.00"
                if column.startswith("VALOR")
                else "#,##0" if column == "QUANTIDADE"
                else "dd/mm/yyyy hh:mm:ss" if column.startswith("DATA_HORA")
                else None
            )
            if number_format:
                for row in range(2, worksheet.max_row + 1):
                    worksheet.cell(row=row, column=index).number_format = number_format

        summary_sheet = writer.sheets["Resumo"]
        for column, width in {"A": 26, "B": 18, "C": 18, "D": 18}.items():
            summary_sheet.column_dimensions[column].width = width
        for row in range(2, len(summary) + 2):
            for column in (3, 4):
                summary_sheet.cell(row=row, column=column).number_format = "#,##0.00"

    tentativas = max(config.SALVAR_TENTATIVAS, 1)
    for tentativa in range(1, tentativas + 1):
        try:
            temporary_file.replace(destino)
            return destino
        except PermissionError:
            if tentativa < tentativas:
                log(
                    f"'{destino.name}' está em uso (tentativa {tentativa}/{tentativas}); "
                    f"nova tentativa em {config.SALVAR_ESPERA_SEG}s"
                )
                time.sleep(config.SALVAR_ESPERA_SEG)

    alternate = destino.with_name(f"{destino.stem}_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    temporary_file.replace(alternate)
    log(
        f"ATENÇÃO: '{destino.name}' continuou em uso; salvei como '{alternate.name}'. "
        f"O BI ainda está lendo a versão anterior."
    )
    return alternate