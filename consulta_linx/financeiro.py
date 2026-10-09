"""Títulos a receber por cliente: nota, pedido, vendedor, devolução e situação de pagamento.

Ligações conferidas na base:
  LANCAMENTOS.TIPO_ORIGEM 'S'/'C' -> ORIGEM = SAIDAS.SAIDA (= NF.COD_OPERACAO)
  LANCAMENTOS.TIPO_ORIGEM 'V'     -> ORIGEM = PEDIDO_VENDA.PEDIDOV (pedido ainda sem nota)
  Devolução = entrada com evento de devolução de venda (EVENTOS.TIPO_ENTRADA = 'V'),
              em que PRODUTOS_EVENTOS.NOTA_REF = número da nota de saída original.

Pedido e devolução são buscados em consultas separadas e cruzados aqui: no Firebird 3
(dialeto 1) a mesma coisa num SELECT só, com subconsulta correlacionada ou LEFT JOIN em
tabela agregada, passa de 10 minutos.
"""

import time
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl.utils import get_column_letter

from . import config
from .database import consultar, consultar_por_ids
from .parquet import montar_schema
from .utils import cod, log, num, substituir_arquivo

SQL_TITULOS = """
SELECT
    L.LANCAMENTO,
    L.TIPO_ORIGEM,
    L.ORIGEM,
    C.COD_CLIENTE,
    C.NOME                 AS CLIENTE,
    L.N_DOCUMENTO          AS TITULO,
    L.DATA_EMISSAO         AS DATA_EMISSAO_TITULO,
    L.DATA_VENCIMENTO,
    L.DATA_PAGAMENTO,
    L.VALOR_INICIAL        AS VALOR_TITULO,
    L.VALOR_PAGO,
    L.EFETUADO,
    NF.NOTA                AS NF,
    NF.DATA                AS DATA_FATURAMENTO,
    S.CLIENTE              AS CLIENTE_NOTA,
    S.FUNCIONARIO          AS FUNCIONARIO_NOTA,
    PVT.COD_PEDIDOV        AS PEDIDO_TITULO,
    PVT.VENDEDOR           AS VENDEDOR_TITULO
FROM LANCAMENTOS L
INNER JOIN CLIENTES C       ON C.CLIENTE = L.COD
LEFT JOIN NF                ON NF.COD_OPERACAO = L.ORIGEM
                           AND NF.TIPO_OPERACAO = 'S'
                           AND L.TIPO_ORIGEM IN ('S', 'C')
LEFT JOIN SAIDAS S          ON S.SAIDA = NF.COD_OPERACAO
LEFT JOIN PEDIDO_VENDA PVT  ON PVT.PEDIDOV = L.ORIGEM
                           AND L.TIPO_ORIGEM = 'V'
WHERE L.TIPO = 'R'
  AND L.GERADOR = 'C'
  AND L.FILIAL = 1
  AND L.DATA_EMISSAO >= ?
  AND C.COD_CLIENTE NOT LIKE '%C%'
  AND L.N_DOCUMENTO NOT STARTING WITH 'VTX-'  -- e-commerce (B2C), pago no ato
"""

SQL_PEDIDOS_POR_SAIDA = """
SELECT DISTINCT PPV.SAIDA, PV.COD_PEDIDOV, PV.VENDEDOR
FROM PRODUTO_PEDIDOV PPV
INNER JOIN PEDIDO_VENDA PV ON PV.PEDIDOV = PPV.PEDIDOV
WHERE PPV.SAIDA IN ({ids})
"""

SQL_DEVOLUCOES = """
SELECT
    PE.NOTA_REF,
    E.CLIENTE,
    PE.NOTA                                AS NF_DEVOLUCAO,
    E.DATA                                 AS DATA_DEVOLUCAO,
    SUM(PE.PRECO_APLICADO * PE.QUANTIDADE) AS VALOR_DEVOLUCAO
FROM ENTRADAS E
INNER JOIN EVENTOS EV          ON EV.EVENTO = E.EVENTO AND EV.TIPO_ENTRADA = 'V'
INNER JOIN PRODUTOS_EVENTOS PE ON PE.COD_OPERACAO = E.ENTRADA AND PE.TIPO_OPERACAO = 'E'
WHERE COALESCE(E.CANCELADA, 'F') = 'F'
  AND PE.NOTA_REF IS NOT NULL
  AND E.DATA >= ?
GROUP BY PE.NOTA_REF, E.CLIENTE, PE.NOTA, E.DATA
"""

SQL_FUNCIONARIOS = "SELECT FUNCIONARIO, NOME FROM FUNCIONARIOS"

COLUNAS_TITULOS = [
    "COD_CLIENTE", "CLIENTE", "COD_VENDEDOR", "VENDEDOR", "NF", "PEDIDO", "TITULO",
    "DATA_EMISSAO_TITULO", "DATA_FATURAMENTO", "DATA_VENCIMENTO", "DATA_PAGAMENTO",
    "VALOR_TITULO", "VALOR_PAGO", "SITUACAO", "DIAS_ATRASO",
    "NF_DEVOLUCAO", "DATA_DEVOLUCAO", "VALOR_DEVOLUCAO",
]
COLUNAS_DATA = [
    "DATA_EMISSAO_TITULO", "DATA_FATURAMENTO", "DATA_VENCIMENTO", "DATA_PAGAMENTO", "DATA_DEVOLUCAO",
]
SCHEMA_TITULOS = montar_schema(
    COLUNAS_TITULOS,
    texto=[
        "COD_CLIENTE", "CLIENTE", "COD_VENDEDOR", "VENDEDOR", "NF", "PEDIDO", "TITULO",
        "SITUACAO", "NF_DEVOLUCAO",
    ],
    data=COLUNAS_DATA,
    inteiro=["DIAS_ATRASO"],
)

PAGO = "PAGO"
VENCIDO = "EM ABERTO - VENCIDO"
A_VENCER = "EM ABERTO - A VENCER"


def extrair(connection, data_inicial: str) -> dict:
    titulos = consultar(connection, SQL_TITULOS, (data_inicial,), "Títulos a receber")
    saidas = titulos.loc[titulos["TIPO_ORIGEM"].isin(["S", "C"]), "ORIGEM"]
    return {
        "titulos": titulos,
        "pedidos": consultar_por_ids(connection, SQL_PEDIDOS_POR_SAIDA, saidas, "Pedidos das notas"),
        "devolucoes": consultar(connection, SQL_DEVOLUCOES, (data_inicial,), "Devoluções"),
        "func": consultar(connection, SQL_FUNCIONARIOS, (), "Funcionários"),
    }


def _juntar_texto(values: pd.Series) -> str:
    return ", ".join(sorted(set(values.dropna())))


def transformar(data: dict, hoje: date | None = None) -> pd.DataFrame:
    started = time.time()
    hoje = pd.Timestamp(hoje or date.today())
    titulos = data["titulos"].copy()
    for column in ["ORIGEM", "NF", "CLIENTE_NOTA", "FUNCIONARIO_NOTA", "VENDEDOR_TITULO"]:
        titulos[column] = cod(titulos[column])
    titulos["SAIDA"] = titulos["ORIGEM"].where(titulos["TIPO_ORIGEM"].isin(["S", "C"]))

    # Pedido e vendedor pela nota; uma nota pode atender mais de um pedido.
    pedidos = data["pedidos"].copy()
    pedidos["SAIDA"] = cod(pedidos["SAIDA"])
    pedidos["VENDEDOR"] = cod(pedidos["VENDEDOR"])
    pedidos = pedidos.groupby("SAIDA", as_index=False).agg(
        PEDIDO_NOTA=("COD_PEDIDOV", _juntar_texto),
        VENDEDOR_PEDIDO=("VENDEDOR", "first"),
    )
    titulos = titulos.merge(pedidos, on="SAIDA", how="left")
    titulos["PEDIDO"] = titulos["PEDIDO_NOTA"].fillna(titulos["PEDIDO_TITULO"])
    titulos["COD_VENDEDOR"] = (
        titulos["VENDEDOR_PEDIDO"]
        .fillna(titulos["VENDEDOR_TITULO"])
        .fillna(titulos["FUNCIONARIO_NOTA"])
    )
    vendedores = data["func"].assign(FUNCIONARIO=lambda df: cod(df["FUNCIONARIO"]))
    titulos["VENDEDOR"] = titulos["COD_VENDEDOR"].map(
        vendedores.drop_duplicates("FUNCIONARIO").set_index("FUNCIONARIO")["NOME"]
    )

    # Devolução: mesma nota de origem e mesmo cliente (o número da nota se repete entre filiais).
    devolucoes = data["devolucoes"].copy()
    devolucoes["NOTA_REF"] = cod(devolucoes["NOTA_REF"])
    devolucoes["CLIENTE"] = cod(devolucoes["CLIENTE"])
    devolucoes["NF_DEVOLUCAO"] = cod(devolucoes["NF_DEVOLUCAO"])
    devolucoes["VALOR_DEVOLUCAO"] = num(devolucoes["VALOR_DEVOLUCAO"])
    devolucoes = devolucoes.groupby(["NOTA_REF", "CLIENTE"], as_index=False).agg(
        NF_DEVOLUCAO=("NF_DEVOLUCAO", _juntar_texto),
        DATA_DEVOLUCAO=("DATA_DEVOLUCAO", "max"),
        VALOR_DEVOLUCAO=("VALOR_DEVOLUCAO", "sum"),
    )
    titulos = titulos.merge(
        devolucoes,
        left_on=["NF", "CLIENTE_NOTA"],
        right_on=["NOTA_REF", "CLIENTE"],
        how="left",
        suffixes=("", "_DEV"),
    )

    for column in COLUNAS_DATA:
        titulos[column] = pd.to_datetime(titulos[column]).dt.normalize()
    for column in ["VALOR_TITULO", "VALOR_PAGO"]:
        titulos[column] = num(titulos[column])
    # Datas digitadas erradas (ex.: ano 2202) não entram no cálculo de atraso.
    data_invalida = titulos["DATA_PAGAMENTO"] > hoje
    if data_invalida.any():
        log(f"Títulos com data de pagamento no futuro (ignorada): {', '.join(titulos.loc[data_invalida, 'TITULO'])}")
        titulos.loc[data_invalida, "DATA_PAGAMENTO"] = pd.NaT

    pago = titulos["EFETUADO"] == "T"
    vencido = ~pago & (titulos["DATA_VENCIMENTO"] < hoje)
    titulos["SITUACAO"] = A_VENCER
    titulos.loc[vencido, "SITUACAO"] = VENCIDO
    titulos.loc[pago, "SITUACAO"] = PAGO
    titulos["DIAS_ATRASO"] = pd.Series(pd.NA, index=titulos.index, dtype="Int64")
    titulos.loc[pago, "DIAS_ATRASO"] = (
        (titulos.loc[pago, "DATA_PAGAMENTO"] - titulos.loc[pago, "DATA_VENCIMENTO"]).dt.days
    )
    titulos.loc[vencido, "DIAS_ATRASO"] = (hoje - titulos.loc[vencido, "DATA_VENCIMENTO"]).dt.days

    result = (
        titulos[COLUNAS_TITULOS]
        .sort_values(["COD_CLIENTE", "DATA_VENCIMENTO", "TITULO"], na_position="last")
        .reset_index(drop=True)
    )
    log(f"Títulos montados: {len(result):,} linhas em {time.time() - started:.1f}s")
    return result


def resumo_por_cliente(titulos: pd.DataFrame) -> pd.DataFrame:
    data = titulos.assign(
        EM_ABERTO=titulos["SITUACAO"] != PAGO,
        VENCIDO=titulos["SITUACAO"] == VENCIDO,
        PAGO_EM_DIA=(titulos["SITUACAO"] == PAGO) & ~(titulos["DIAS_ATRASO"] > 0),
        PAGO_EM_ATRASO=(titulos["SITUACAO"] == PAGO) & (titulos["DIAS_ATRASO"] > 0),
    )
    data["VALOR_EM_ABERTO"] = data["VALOR_TITULO"].where(data["EM_ABERTO"], 0)
    data["VALOR_VENCIDO"] = data["VALOR_TITULO"].where(data["VENCIDO"], 0)
    summary = data.groupby(["COD_CLIENTE", "CLIENTE"], as_index=False).agg(
        TITULOS=("TITULO", "size"),
        EM_ABERTO=("EM_ABERTO", "sum"),
        VENCIDOS=("VENCIDO", "sum"),
        PAGOS_EM_DIA=("PAGO_EM_DIA", "sum"),
        PAGOS_EM_ATRASO=("PAGO_EM_ATRASO", "sum"),
        VALOR_EM_ABERTO=("VALOR_EM_ABERTO", "sum"),
        VALOR_VENCIDO=("VALOR_VENCIDO", "sum"),
        ULTIMO_PAGAMENTO=("DATA_PAGAMENTO", "max"),
    )
    pagos = summary["PAGOS_EM_DIA"] + summary["PAGOS_EM_ATRASO"]
    summary["PERCENTUAL_ATRASO"] = (summary["PAGOS_EM_ATRASO"] / pagos.where(pagos > 0)).round(4)
    # A devolução se repete em cada parcela da nota: soma uma vez por nota.
    devolvido = (
        titulos.dropna(subset=["NF_DEVOLUCAO"])
        .drop_duplicates(["COD_CLIENTE", "NF"])
        .groupby("COD_CLIENTE")
        .agg(NOTAS_COM_DEVOLUCAO=("NF", "size"), VALOR_DEVOLVIDO=("VALOR_DEVOLUCAO", "sum"))
    )
    summary = summary.merge(devolvido, on="COD_CLIENTE", how="left")
    summary["NOTAS_COM_DEVOLUCAO"] = summary["NOTAS_COM_DEVOLUCAO"].fillna(0).astype(int)
    summary["VALOR_DEVOLVIDO"] = summary["VALOR_DEVOLVIDO"].fillna(0.0)
    return summary.sort_values(["VALOR_VENCIDO", "PAGOS_EM_ATRASO"], ascending=False).reset_index(drop=True)


def salvar_excel(titulos: pd.DataFrame, resumo: pd.DataFrame, destino: Path) -> Path:
    destino.parent.mkdir(parents=True, exist_ok=True)
    temporary_file = destino.with_name(destino.stem + ".tmp.xlsx")
    with pd.ExcelWriter(
        temporary_file, engine="openpyxl", datetime_format="dd/mm/yyyy", date_format="dd/mm/yyyy"
    ) as writer:
        for sheet, data in {"Titulos": titulos, "Resumo por cliente": resumo}.items():
            data.to_excel(writer, sheet_name=sheet, index=False)
            worksheet = writer.sheets[sheet]
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = f"A1:{get_column_letter(data.shape[1])}{len(data) + 1}"
            for index, column in enumerate(data.columns, start=1):
                width = 40 if column in ("CLIENTE", "VENDEDOR") else max(12, len(column) + 2)
                worksheet.column_dimensions[get_column_letter(index)].width = width
                number_format = (
                    "0.0%" if column == "PERCENTUAL_ATRASO"
                    else "#,##0.00" if column.startswith("VALOR")
                    else None
                )
                if number_format:
                    for row in range(2, len(data) + 2):
                        worksheet.cell(row=row, column=index).number_format = number_format
        pd.DataFrame(
            {"PARÂMETRO": ["Gerado em", "Títulos emitidos a partir de"],
             "VALOR": [datetime.now().strftime("%d/%m/%Y %H:%M"),
                       pd.Timestamp(config.FINANCEIRO_DATA_INICIAL).strftime("%d/%m/%Y")]}
        ).to_excel(writer, sheet_name="Info", index=False)
    return substituir_arquivo(temporary_file, destino)
